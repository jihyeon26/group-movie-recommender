"""Build the current scientific report from frozen final-holdout artifacts."""

from __future__ import annotations

import json
from pathlib import Path
import sys

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = PROJECT_ROOT / "outputs/final_holdout_experiment"
REPORT_PATH = PROJECT_ROOT / "reports/group_movie_recommender_report.docx"
ASSET_DIR = PROJECT_ROOT / "reports/assets"

NAVY = "183B56"
TEAL = "15807A"
LIGHT_BLUE = "EAF1F5"
LIGHT_TEAL = "E7F3F1"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "66737F"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths_cm: list[float]) -> None:
    table.autofit = False
    total_twips = int(sum(widths_cm) / 2.54 * 1440)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width / 2.54 * 1440)))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_table(document, headers: list[str], rows: list[list[str]], widths_cm: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(8.2)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = str(value)
            if row_index % 2:
                set_cell_shading(cells[column_index], LIGHT_GRAY)
            paragraph = cells[column_index].paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            for run in paragraph.runs:
                run.font.size = Pt(8.2)
    set_table_geometry(table, widths_cm)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, instruction_node, separate])
    if placeholder:
        paragraph.add_run(placeholder)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, " PAGE ", "1")
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_toc(document) -> None:
    paragraph = document.add_paragraph()
    add_field(paragraph, ' TOC \\o "1-2" \\h \\z \\u ', "Right-click and update field")


def add_caption(document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="Caption Compact")
    paragraph.paragraph_format.keep_with_next = True


def add_callout(document, label: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_TEAL)
    set_cell_margins(cell, top=150, start=170, bottom=150, end=170)
    paragraph = cell.paragraphs[0]
    lead = paragraph.add_run(f"{label}  ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(TEAL)
    paragraph.add_run(text)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [16.8])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(document, text: str, *, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    paragraph = document.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_after = Pt(2)


def set_document_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(34, 45, 56)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color, before, after in (
        ("Title", 27, NAVY, 0, 8),
        ("Subtitle", 12, MID_GRAY, 0, 8),
        ("Heading 1", 16, NAVY, 11, 5),
        ("Heading 2", 11.5, TEAL, 8, 3),
        ("Heading 3", 10, NAVY, 6, 2),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(9.2)
        styles[name].paragraph_format.left_indent = Cm(0.55)
        styles[name].paragraph_format.first_line_indent = Cm(-0.3)
    if "Caption Compact" not in styles:
        caption = styles.add_style("Caption Compact", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption Compact"]
    caption.font.name = "Arial"
    caption.font.size = Pt(8.2)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(3)


def configure_sections(document: Document) -> None:
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.55)
        section.left_margin = Cm(1.85)
        section.right_margin = Cm(1.85)
        section.header_distance = Cm(0.6)
        section.footer_distance = Cm(0.65)
        section.header.paragraphs[0].text = ""
        add_page_number(section)


def make_result_figure(comparison: pd.DataFrame, output_path: Path) -> None:
    labels = ["Popularity", "LightGCN", "ItemKNN"]
    values = comparison.set_index("method").loc[
        ["popularity", "lightgcn_average", "item_knn_average"]
    ]
    metrics = ["meanMinimumNDCG@10", "meanAverageNDCG@10", "catalogueCoverage@10"]
    titles = ["Minimum-member NDCG@10", "Average NDCG@10", "Catalogue coverage@10"]
    colors = ["#8A98A8", "#2D74B8", "#15807A"]
    image = Image.new("RGB", (2040, 600), "white")
    draw = ImageDraw.Draw(image)
    font_dir = Path("C:/Windows/Fonts")
    title_font = ImageFont.truetype(str(font_dir / "arialbd.ttf"), 29)
    label_font = ImageFont.truetype(str(font_dir / "arial.ttf"), 22)
    value_font = ImageFont.truetype(str(font_dir / "arial.ttf"), 20)
    panel_width = 640
    for panel, (metric, title) in enumerate(zip(metrics, titles)):
        left = 20 + panel * 675
        right = left + panel_width
        top, bottom = 80, 500
        numbers = values[metric].to_numpy(dtype=float)
        maximum = max(numbers) * 1.25
        title_box = draw.textbbox((0, 0), title, font=title_font)
        draw.text(
            (left + (panel_width - (title_box[2] - title_box[0])) / 2, 20),
            title,
            fill="#183B56",
            font=title_font,
        )
        draw.line((left + 40, bottom, right - 20, bottom), fill="#C6CED5", width=2)
        bar_width, gap = 125, 48
        start_x = left + 60
        for index, (label, number, color) in enumerate(zip(labels, numbers, colors)):
            x0 = start_x + index * (bar_width + gap)
            height = int((number / maximum) * (bottom - top))
            y0 = bottom - height
            draw.rounded_rectangle((x0, y0, x0 + bar_width, bottom), radius=8, fill=color)
            value = f"{number:.4f}"
            value_box = draw.textbbox((0, 0), value, font=value_font)
            draw.text(
                (x0 + (bar_width - (value_box[2] - value_box[0])) / 2, y0 - 28),
                value,
                fill="#263746",
                font=value_font,
            )
            label_box = draw.textbbox((0, 0), label, font=label_font)
            draw.text(
                (x0 + (bar_width - (label_box[2] - label_box[0])) / 2, bottom + 12),
                label,
                fill="#3F4C58",
                font=label_font,
            )
    image.save(output_path, dpi=(220, 220))


def build_report() -> Path:
    selection = json.loads((OUTPUT_ROOT / "selection.json").read_text(encoding="utf-8"))
    protocol = json.loads((OUTPUT_ROOT / "protocol.json").read_text(encoding="utf-8"))
    report = json.loads(
        (OUTPUT_ROOT / "holdout/holdout_report.json").read_text(encoding="utf-8")
    )
    comparison = pd.read_csv(OUTPUT_ROOT / "holdout/comparison.csv")
    subgroup = pd.read_csv(OUTPUT_ROOT / "holdout/similarity_subgroups.csv")
    bootstrap = pd.DataFrame(report["paired_bootstrap"])
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    figure_path = ASSET_DIR / "final_holdout_performance.png"
    make_result_figure(comparison, figure_path)

    document = Document()
    document.core_properties.author = "Jihyeon Choung"
    document.core_properties.last_modified_by = "Jihyeon Choung"
    document.core_properties.title = "Shared Movie Recommendation for Randomly Paired Users"
    document.core_properties.subject = "Recommender Systems scientific report"
    set_document_styles(document)
    configure_sections(document)
    settings = document.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    # Page 1: cover.
    document.add_paragraph().paragraph_format.space_after = Pt(74)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Shared Movie Recommendation\nfor Randomly Paired Users")
    subtitle = document.add_paragraph(
        "A Temporal Comparison of Popularity, LightGCN, and ItemKNN",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph().paragraph_format.space_after = Pt(58)
    for value in ("Jihyeon Choung", "MSc Applied Information and Data Science", "1 September 2026"):
        paragraph = document.add_paragraph(value)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(7)
    document.add_page_break()

    # Page 2: linked TOC.
    document.add_heading("Contents", level=0)
    add_toc(document)
    document.add_page_break()

    # Page 3: abstract and framing.
    document.add_heading("Abstract", level=1)
    document.add_paragraph(
        "Choosing one movie for two people is a group recommendation problem: a list can "
        "perform well on average while serving one member poorly. This study compares global "
        "popularity, LightGCN, and item-based nearest neighbours (ItemKNN) on MovieLens 32M. "
        "Users were paired uniformly at random without replacement from a pre-test eligible "
        "cohort; taste similarity did not determine membership. Models were selected on 300 "
        "development pairs and evaluated once on 500 user-disjoint holdout pairs. Of these, 195 "
        "had observable positive test items for both members. ItemKNN achieved minimum-member "
        "NDCG@10 of 0.0335 versus 0.0260 for popularity, an absolute paired improvement of "
        "+0.0075 (95% bootstrap CI [0.0005, 0.0148]). LightGCN did not outperform popularity. "
        "The results support a simple neighbourhood model for this controlled task, but sparse "
        "labels and synthetic pairs limit conclusions about real joint satisfaction."
    )
    document.add_paragraph(
        "Keywords: group recommendation; MovieLens 32M; LightGCN; ItemKNN; temporal evaluation"
    )
    document.add_heading("1. Problem Framing and Motivation", level=1)
    document.add_paragraph(
        "The target user is a pair deciding what to watch on a shared screen. The system returns "
        "one Top-10 list, not two independent lists. Because average utility can conceal a poor "
        "experience for one person, the primary outcome is minimum-member NDCG@10, averaged "
        "across pairs. Average NDCG@10 captures overall utility and catalogue coverage measures "
        "how broadly each method uses the available catalogue."
    )
    add_callout(
        document,
        "Research question",
        "Under a temporal warm-catalogue evaluation on randomly paired, member-disjoint "
        "MovieLens users, do LightGCN or ItemKNN improve minimum-member NDCG@10 over global "
        "popularity, and what trade-offs do they create in average NDCG@10 and catalogue "
        "coverage@10?",
    )
    document.add_heading("1.1 Contribution", level=2)
    add_bullet(document, "A reproducible random-pair protocol that does not select groups by taste similarity.")
    add_bullet(document, "A three-model comparison on identical histories, candidates, pairs, and labels.")
    add_bullet(document, "Paired uncertainty estimates and a prespecified descriptive similarity analysis.")
    document.add_page_break()

    # Page 4: data and cohort construction.
    document.add_heading("2. Dataset and Feedback", level=1)
    document.add_paragraph(
        "MovieLens 32M contains 32,000,204 explicit ratings from 200,948 users for 87,585 "
        "movies [1]. Ratings range from 0.5 to 5.0. A rating of at least 4.0 is treated as a "
        "positive implicit interaction for graph construction and relevance labels. Lower "
        "ratings remain available for descriptive analysis but are not interpreted as observed "
        "negative feedback."
    )
    add_caption(document, "Table 1. Temporal split and its role in the experiment.")
    add_table(
        document,
        ["Period", "Timestamp rule", "Use"],
        [
            ["Train", "Before 2019-01-01", "Positive graph, popularity, ItemKNN, pair descriptions"],
            ["Development", "2019 calendar year", "LightGCN checkpoint selection only"],
            ["Holdout", "2020-01-01 onward", "Single final comparison after freezing models"],
        ],
        [2.4, 4.2, 10.2],
    )
    document.add_heading("2.1 Random pair construction", level=2)
    document.add_paragraph(
        "Eligibility used only pre-test information: at least 100 train ratings, 20 positive "
        "train ratings, and three positive development ratings. The 1,003 users appearing in "
        "earlier exploratory pair cohorts were excluded. From 2,574 remaining eligible users, "
        "1,600 users were sampled uniformly without replacement with seed 20260828 and paired "
        "adjacently. The first 300 pairs formed the development cohort and the remaining 500 "
        "formed the holdout cohort. No user appears in more than one pair or in both cohorts."
    )
    add_callout(
        document,
        "Design decision",
        "Genre and rating similarities describe the random pairs but never determine who is "
        "paired. This represents an application where a system receives an existing pair rather "
        "than choosing only unusually incompatible users.",
    )
    document.add_heading("2.2 Controlled graph", level=2)
    graph = protocol["graph"]
    document.add_paragraph(
        f"All 1,600 pair members were retained in a controlled graph with {graph['users']:,} "
        f"users, {graph['movies']:,} warm movies, and {graph['edges']:,} positive train edges. "
        "Context users were sampled reproducibly. This graph is substantially smaller than the "
        "full dataset, making all three methods feasible on the same hardware while preserving "
        "a sizeable catalogue."
    )
    document.add_page_break()

    # Page 5: task and methods.
    document.add_heading("3. Recommendation Task and Candidate Set", level=1)
    document.add_paragraph(
        "For each pair (A, B), each model scores the warm graph catalogue. Movies rated by "
        "either member before the target period are removed, and the remaining movies form the "
        "shared candidate set. Individual scores are converted to within-user percentile scores "
        "before averaging, preventing one model member's score scale from dominating. The final "
        "output is one shared Top-10 list. A positive target is an unseen movie rated at least "
        "4.0 during the evaluated period."
    )
    document.add_heading("4. Compared Methods", level=1)
    document.add_heading("4.1 Popularity baseline", level=2)
    document.add_paragraph(
        "Global popularity ranks movies by positive train-interaction count after removing the "
        "pair's seen union. It is non-personalized and provides the minimum comparison needed to "
        "justify model complexity."
    )
    document.add_heading("4.2 ItemKNN", level=2)
    document.add_paragraph(
        "ItemKNN computes cosine similarity between binary positive-interaction item columns. "
        "Each user's score is the sum of similarities from previously liked movies. The model "
        "uses 100 neighbours and shrinkage 10.0, then averages the two normalized member scores. "
        "This is a personalized classical method with transparent local evidence."
    )
    document.add_heading("4.3 LightGCN", level=2)
    light = selection["lightgcn"]
    document.add_paragraph(
        "LightGCN learns user and item embeddings by propagating signals over the bipartite "
        "positive-interaction graph and optimizing Bayesian Personalized Ranking loss [2,3]. "
        "The configuration uses 32-dimensional embeddings, two graph layers, batch size 512, "
        "learning rate 0.02, L2 weight 0.0001, and seed 20260828. Validation ran every 100 steps "
        f"with patience five. Training stopped at step {light['stopped_step']:,} and restored the "
        f"interior best checkpoint at step {light['best_step']:,} from a maximum budget of "
        f"{light['config']['steps']:,}. The two normalized member scores are averaged."
    )
    add_caption(document, "Table 2. Model roles in the controlled comparison.")
    add_table(
        document,
        ["Method", "Personalized", "Primary purpose"],
        [
            ["Popularity", "No", "Simple baseline"],
            ["ItemKNN", "Yes", "Classical neighbourhood comparison"],
            ["LightGCN", "Yes", "Advanced graph recommendation candidate"],
        ],
        [3.3, 2.6, 10.9],
    )
    document.add_page_break()

    # Page 6: evaluation protocol.
    document.add_heading("5. Evaluation Protocol", level=1)
    document.add_paragraph(
        "The experiment has three irreversible phases. Preparation freezes eligible users, "
        "random pairs, graph edges, and file hashes without loading 2020+ feedback. Training "
        "selects the LightGCN checkpoint using only development users and 2019 labels, then "
        "freezes LightGCN and ItemKNN artifacts. The test phase verifies all hashes and evaluates "
        "the 500 holdout pairs once. Earlier exploratory users were excluded, although the same "
        "calendar period had been inspected for those different users."
    )
    document.add_heading("5.1 Metrics", level=2)
    add_caption(document, "Table 3. Prespecified metrics.")
    add_table(
        document,
        ["Metric", "Role", "Interpretation"],
        [
            ["Minimum-member NDCG@10", "Primary", "Rank quality received by the worse-off member"],
            ["Average NDCG@10", "Secondary", "Mean ranked relevance across both members"],
            ["Minimum / average Recall@10", "Secondary", "Recovery of positive future items"],
            ["Catalogue coverage@10", "Beyond accuracy", "Share of warm items recommended at least once"],
            ["Two-sided hit rate", "Diagnostic", "Pairs for which both members receive at least one hit"],
        ],
        [4.5, 3.2, 9.1],
    )
    document.add_heading("5.2 Evaluable pairs and uncertainty", level=2)
    document.add_paragraph(
        "Pair-level NDCG and Recall require at least one eligible positive target for each member. "
        "This was true for 195 of 500 frozen holdout pairs (39.0%); the other 305 were retained in "
        "the cohort record but excluded from two-sided metric means. Paired nonparametric "
        "bootstrap intervals use the 195 common evaluable pairs, 10,000 resamples, and seed "
        "20260828. The 16,923-movie catalogue and candidate filtering are identical across models."
    )
    add_callout(
        document,
        "Interpretation boundary",
        "The bootstrap quantifies uncertainty for observable two-sided pairs. It does not correct "
        "selection toward users who remain active enough to leave ratings in the test period.",
    )
    document.add_page_break()

    # Page 7: main result.
    document.add_heading("6. Final Holdout Results", level=1)
    document.add_paragraph(
        "ItemKNN is the only personalized model that outperforms popularity on both main NDCG "
        "outcomes. It improves minimum-member NDCG@10 from 0.0260 to 0.0335 and average "
        "NDCG@10 from 0.0959 to 0.1100. LightGCN provides the highest catalogue coverage but "
        "does not improve relevance. Absolute metric values remain low, which is expected under "
        "a shared unseen-list task with sparse future ratings."
    )
    rows = []
    display_names = {
        "popularity": "Popularity",
        "lightgcn_average": "LightGCN",
        "item_knn_average": "ItemKNN",
    }
    for _, row in comparison.iterrows():
        rows.append([
            display_names[row["method"]],
            f"{row['meanMinimumNDCG@10']:.4f}",
            f"{row['meanAverageNDCG@10']:.4f}",
            f"{row['meanMinimumRecall@10']:.4f}",
            f"{row['meanAverageRecall@10']:.4f}",
            f"{row['catalogueCoverage@10']:.4f}",
        ])
    add_caption(document, "Table 4. Frozen holdout comparison on 195 common evaluable random pairs.")
    add_table(
        document,
        ["Method", "Min NDCG", "Avg NDCG", "Min Recall", "Avg Recall", "Coverage"],
        rows,
        [3.3, 2.7, 2.7, 2.7, 2.7, 2.7],
    )
    result_figure = document.add_picture(str(figure_path), width=Inches(6.45))
    result_figure._inline.docPr.set(
        "descr",
        "Bar charts comparing minimum-member NDCG, average NDCG, and catalogue "
        "coverage for Popularity, LightGCN, and ItemKNN.",
    )
    result_figure._inline.docPr.set("title", "Final holdout performance comparison")
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(document, "Figure 1. Relevance and catalogue coverage on the final holdout cohort.")
    relevant_bootstrap = bootstrap[
        ((bootstrap["methodA"] == "item_knn_average") & (bootstrap["methodB"] == "popularity"))
        | ((bootstrap["methodA"] == "lightgcn_average") & (bootstrap["methodB"] == "popularity"))
    ]
    bootstrap_rows = []
    for row in relevant_bootstrap.itertuples(index=False):
        method = "ItemKNN" if row.methodA == "item_knn_average" else "LightGCN"
        metric = "Minimum NDCG" if row.metric.startswith("minimum") else "Average NDCG"
        bootstrap_rows.append([
            f"{method} - Popularity",
            metric,
            f"{row.meanDifferenceAminusB:+.5f}",
            f"[{row.ci95Lower:+.5f}, {row.ci95Upper:+.5f}]",
        ])
    add_caption(document, "Table 5. Paired bootstrap differences; positive values favour the personalized method.")
    add_table(document, ["Comparison", "Metric", "Difference", "95% CI"], bootstrap_rows, [5.2, 3.6, 3.4, 4.6])
    document.add_page_break()

    # Page 8: subgroup analysis and discussion.
    document.add_heading("6.1 Descriptive Taste-Similarity Analysis", level=2)
    rule = selection["similarity_subgroup_rule"]
    document.add_paragraph(
        "Taste similarity was not used for pairing. For interpretation only, development-pair "
        "genre similarity fixed two cut-offs (0.0736 and 0.3491) before holdout evaluation. "
        "Applying them yielded 54 high-conflict, 63 mixed, and 78 similar evaluable holdout "
        "pairs. These subgroup estimates were not used for model selection and do not include "
        "separate confidence intervals."
    )
    subgroup_rows = []
    band_names = {"high_conflict": "High conflict", "mixed": "Mixed", "similar": "Similar"}
    for band in ("high_conflict", "mixed", "similar"):
        values = subgroup[subgroup["similarityBand"] == band].set_index("method")
        for method in ("popularity", "lightgcn_average", "item_knn_average"):
            row = values.loc[method]
            subgroup_rows.append([
                band_names[band],
                display_names[method],
                str(int(row["evaluatedPairs"])),
                f"{row['meanMinimumNDCG']:.4f}",
                f"{row['meanAverageNDCG']:.4f}",
            ])
    add_caption(document, "Table 6. Descriptive results by pre-frozen genre-similarity band.")
    add_table(
        document,
        ["Band", "Method", "Pairs", "Min NDCG", "Avg NDCG"],
        subgroup_rows,
        [3.3, 3.5, 2.2, 3.9, 3.9],
    )
    document.add_paragraph(
        "ItemKNN's overall advantage is concentrated in mixed and similar pairs. In the "
        "high-conflict subgroup, it improves average NDCG over popularity (0.0746 versus 0.0643) "
        "but has lower minimum-member NDCG (0.0069 versus 0.0109). Therefore the main result "
        "does not support claiming that ItemKNN protects the worse-off member specifically when "
        "tastes are most dissimilar."
    )
    document.add_heading("7. Discussion, Risks, and Responsible Use", level=1)
    add_bullet(document, "Synthetic pairs approximate a joint choice; separate future ratings are not observed co-satisfaction.")
    add_bullet(document, "Only 195 of 500 pairs have two-sided observable targets, creating activity-selection risk.")
    add_bullet(document, "The 5,000-user graph is controlled rather than full-scale MovieLens 32M training.")
    add_bullet(document, "Missing ratings are unknown, not confirmed dislikes; offline hits underestimate possible enjoyment.")
    add_bullet(document, "One pair seed and one model seed limit robustness across random realizations.")
    add_bullet(document, "Popularity and neighbourhood methods may reproduce exposure bias toward established movies.")
    document.add_paragraph(
        "A deployed system should present the shared list as a compromise and allow both members "
        "to veto or adjust it. A user study should measure joint acceptance, perceived fairness, "
        "and whether one member repeatedly yields to the other."
    )
    document.add_page_break()

    # Page 9: reproducibility, contribution, and conclusion.
    document.add_heading("8. Reproducibility", level=1)
    document.add_paragraph(
        "Repository: github.com/jihyeon26/group-movie-recommender. Data and generated outputs "
        "are excluded from version control, while configurations, phased pipeline code, and tests "
        "are committed. Preparation records SHA-256 hashes for frozen pairs, graph edges, and "
        "configuration files; the test phase refuses altered inputs or a LightGCN checkpoint "
        "selected at the training-budget boundary. All 74 unit tests pass in the current workspace."
    )
    add_caption(document, "Table 7. Principal reproducibility artifacts.")
    add_table(
        document,
        ["Artifact", "Purpose"],
        [
            ["configs/final_holdout_experiment.json", "Cohort, model, and bootstrap parameters"],
            ["scripts/run_final_holdout_experiment.py", "Prepare, train, and one-time test entry point"],
            ["outputs/.../protocol.json", "Frozen cohort rules, graph counts, and input hashes"],
            ["outputs/.../selection.json", "Development-selected checkpoint and model hashes"],
            ["outputs/.../holdout/holdout_report.json", "Final metrics, intervals, risks, and subgroup summary"],
        ],
        [7.1, 9.7],
    )
    document.add_heading("9. Conclusion", level=1)
    document.add_paragraph(
        "Under the frozen random-pair protocol, ItemKNN outperforms popularity on both minimum-"
        "member and average NDCG@10, with paired confidence intervals above zero. LightGCN does "
        "not outperform either simpler method, although it recommends from a broader part of the "
        "catalogue. The scientific conclusion is therefore not that a graph neural model is best, "
        "but that a well-matched classical personalized baseline is strongest for this task and "
        "data regime. The principal remaining uncertainty is whether the result transfers from "
        "synthetic, active MovieLens pairs to real joint decisions."
    )
    document.add_heading("10. Contribution and Use of AI Tools", level=1)
    document.add_heading("10.1 Author contribution", level=2)
    document.add_paragraph(
        "Jihyeon Choung: problem formulation, data validation, temporal protocol, pair design, "
        "model implementation, experiment execution, result interpretation, and report review."
    )
    document.add_heading("10.2 Use of AI tools", level=2)
    document.add_paragraph(
        "OpenAI Codex was used as a coding and writing assistant for repository organization, "
        "draft implementations, debugging, experiment planning, document generation, and language "
        "editing. The author reviewed the code and claims, executed experiments locally, inspected "
        "machine-readable outputs, and remains responsible for the final work. AI was not used to "
        "generate ratings or replace MovieLens observations."
    )
    document.add_page_break()

    # Page 10: references.
    document.add_heading("References", level=1)
    references = [
        "[1] GroupLens Research. MovieLens 32M Dataset README. University of Minnesota, 2024. https://grouplens.org/datasets/movielens/32m/",
        "[2] X. He, K. Deng, X. Wang, Y. Li, Y. Zhang, and M. Wang. LightGCN: Simplifying and Powering Graph Convolution Network for Recommendation. SIGIR, 2020.",
        "[3] S. Rendle, C. Freudenthaler, Z. Gantner, and L. Schmidt-Thieme. BPR: Bayesian Personalized Ranking from Implicit Feedback. UAI, 2009.",
        "[4] B. Sarwar, G. Karypis, J. Konstan, and J. Riedl. Item-Based Collaborative Filtering Recommendation Algorithms. WWW, 2001.",
        "[5] K. Jarvelin and J. Kekalainen. Cumulated Gain-Based Evaluation of IR Techniques. ACM TOIS 20(4), 2002.",
        "[6] J. Masthoff. Group Recommender Systems: Combining Individual Models. In Recommender Systems Handbook, 2011.",
    ]
    for reference in references:
        paragraph = document.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Cm(0.55)
        paragraph.paragraph_format.first_line_indent = Cm(-0.55)
        paragraph.paragraph_format.space_after = Pt(7)
    document.add_paragraph(
        "All reported numerical results were generated locally from MovieLens 32M using the "
        "committed experiment configuration. Generated data and model artifacts are not included "
        "in the repository."
    )

    # Page 11: reflection, explicitly outside the ten-page report limit.
    reflection_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_sections(document)
    document.add_heading("Learning Reflection", level=1)
    document.add_paragraph(
        "This project changed my view of recommender systems from a model-building task into an "
        "evaluation-design task. The most important decision was changing pair construction. I "
        "initially focused on users from the bottom quartile of taste similarity because they made "
        "the conflict problem look clear. However, selecting only unusually incompatible pairs "
        "made the application less natural and tied the conclusion to an arbitrary threshold. "
        "Random, member-disjoint pairing produced a clearer question: given a pair that already "
        "exists, which model produces the better shared list? Taste similarity then became an "
        "analysis variable rather than a gatekeeper."
    )
    document.add_paragraph(
        "I also learned why a popularity baseline cannot be treated as a formality. LightGCN was "
        "the most advanced model I implemented, but it did not beat popularity on the final "
        "holdout relevance metrics. ItemKNN, a simpler neighbourhood method, performed best. This "
        "result corrected my initial assumption that a graph neural recommender would be stronger "
        "because it was more sophisticated. The appropriate comparison line is therefore not model "
        "complexity; it is performance under the same pairs, histories, candidates, and labels."
    )
    document.add_paragraph(
        "The training process made checkpoint selection concrete. An earlier experiment ended at "
        "1,000 steps, which was only a computation limit and not a scientific stopping rule. In the "
        "final experiment I allowed 2,500 steps, evaluated every 100, and used patience five. The "
        "best checkpoint occurred at step 800 and training stopped at 1,300, so the reported model "
        "is an interior validation choice rather than a boundary artifact. I can now explain why "
        "the model stopped and which data selected it."
    )
    document.add_paragraph(
        "Finally, the 195 evaluable pairs out of 500 showed that a larger nominal cohort does not "
        "automatically create more evidence. Both members need observable future positives for the "
        "minimum-member metric, and missing ratings are not dislikes. I would next repeat the random "
        "pairing and training seeds and, if possible, run a small study in which real pairs jointly "
        "accept or reject recommendations. That would connect offline ranking metrics to the social "
        "decision the system is intended to support."
    )

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(REPORT_PATH)
    print(REPORT_PATH)
    return REPORT_PATH


if __name__ == "__main__":
    try:
        build_report()
    except FileNotFoundError as error:
        sys.exit(f"Missing frozen experiment artifact: {error}")
